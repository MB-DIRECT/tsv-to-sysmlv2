#!/usr/bin/env python3
"""
tsv_to_sysmlv2.py -- deterministic SysML 1.7 -> SysML v2 converter.

Part of: Metamodel Scope Using Model-Based Decision on Intended-Use & Requirements Engineering Categorization Tool (MB-DIRECT), Doctor of Engineering praxis, The George Washington
         University, School of Engineering and Applied Science, 2026.
Author:  Evan Matthews (ORCID 0009-0005-2265-0754)
License: MIT (see LICENSE in the repository root)
AI assistance: developed with assistance from Anthropic's Claude (Claude Code);
         specified, reviewed, tested, and validated by the author in
         accordance with the AI Policy for George Washington University Online Engineering Programs.

Input:  the 33-column TSV exported from a SysML 1 model (MSOSA 2024 was used). Attributes in Column title order:
Phase
SysML 1.7 Element Type
Name
Owner Qualified Name
Applied Stereotype
Qualified Name
Type Qualified Name
Multiplicity
Direction
Source/Client Qualified Name
Src End Part With Port QN
Target/Supplier Qualified Name
Tgt End Part With Port QN
Item Conveyed Qualified Name
Called Behavior Qualified Name
Source Node Qualified Name
Target Node Qualified Name
Trigger
Guard Expression
Signal Qualified Name
Element ID
Diagram Type
General/Parent QN
CE0 Role QN
CE1 Role QN
State Entry Behavior QN
State Exit Behavior QN
State Do Activity QN
Transition Effect QN
Pseudostate Kind
Partition Represents QN
Association End 0
Association End 1


Output: 1) a SysML v2 textual model (.sysml or .txt)
        2) a migration report .xlsx in the same layout as the CATIA plugin's
           MigrationReportOf<project>.xlsx:
             - 'Not Migrated Elements'   (Element MetaType | Element Name |
                Qualified Name | Element Url | Element ID)
             - 'Not Migrated Properties' (Element Name | Element MetaType |
                Property Name | Qualified Name | Element Url, with
                continuation rows per extra property)
             - 'Not Migrated Diagrams'   (extra sheet, same column style)

Usage:
  python tsv_to_sysmlv2.py Export.tsv
  python tsv_to_sysmlv2.py Export.tsv -o out.sysml --dialect pilot
  python tsv_to_sysmlv2.py Export.tsv --dialect msosa --report MyReport.xlsx
  python tsv_to_sysmlv2.py Export.tsv --validate            (pilot kernel check)

Dialects:
  pilot (default) - loads clean in the SysML v2 Pilot Implementation
                    (Anaconda Jupyter 'sysml' kernel): named state
                    entry/do/exit actions, conservative trigger/guard
                    comments, requirement usages.
  msosa           - CATIA Magic / MSOSA 2026x conventions: colon-typed state
                    actions allowed; otherwise identical mappings.

Element-type mapping follows the validated June 2026 hand translation:
  Block->part def, InterfaceBlock->port def, ValueType->attribute def,
  Signal->item def, ConstraintBlock->constraint def, Requirement->requirement
  usage, DeriveReqt->named dependency, UseCase->use case def,
  Activity->action def, StateMachine->state def, FlowProperty->in/out item,
  ItemFlow->flow (msosa dialect; pilot needs port dot-paths -> reported),
  BindingConnector->named binding/bind, Connector->connect,
  Association->connection def. Activity internals (2026-08 upgrade):
  ControlFlow->named succession, ObjectFlow->named succession flow (pin
  dot-paths), Merge/Decision/Fork/Join Node->merge/decide/fork/join,
  Initial/ActivityFinal Node->named action + start/done succession,
  AcceptEventAction->action{accept payload:Signal}, SendSignalAction->named
  action (signal/target not in TSV -> property row), Input/OutputPin->in/out
  parameters. State machines: a NAMED Region stays a nested composite state
  (v1.7 QNs survive), an initial Pseudostate keeps its name as a state routed
  from 'entry', and the initial Transition keeps its name.
Transition-owned effect Behaviors (e.g. an internal transition typed by an
Activity) are hoisted into the enclosing scope as full action defs and wired
back via the transition's 'do' clause; the hoisted definition's qualified
name shifts one level up (noted in the properties report).
Anything without a clean v2 home becomes a '// TODO-MIGRATE …' stub AND a row
in the report; named elements absorbed by a parent construct (triggers,
signal events) also get a report row, and a structural safety net turns any
element no handler reached into a report row (nothing is silently dropped).

Dependencies: openpyxl (report only). --validate additionally needs the
SysML v2 Pilot Implementation Jupyter kernel ('sysml') reachable through a
jupyter executable on the PATH, or one given with --jupyter.
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse

# ---------------------------------------------------------------- constants

HDR_TYPE = "SysML 1.7 Element Type"
# jupyter is resolved from the PATH; pass --jupyter if the SysML v2 Pilot
# kernel lives in another environment (for example a separate Anaconda install).
JUPYTER_DEFAULT = shutil.which("jupyter") or "jupyter"

PRIMITIVES = {"Real": "ScalarValues::Real", "Integer": "ScalarValues::Integer",
              "String": "ScalarValues::String", "Boolean": "ScalarValues::Boolean",
              "Natural": "ScalarValues::Natural", "Rational": "ScalarValues::Rational",
              "Number": "ScalarValues::Number"}

RESERVED = {
    "about", "abstract", "accept", "action", "actor", "after", "alias", "all",
    "allocate", "allocation", "analysis", "and", "as", "assert", "assign",
    "assume", "at", "attribute", "bind", "binding", "by", "calc", "case",
    "comment", "concern", "connect", "connection", "constraint", "decide",
    "def", "default", "defined", "dependency", "derived", "do", "doc", "else",
    "end", "entry", "enum", "event", "exhibit", "exit", "expose", "filter",
    "first", "flow", "for", "fork", "frame", "from", "hastype", "if", "import",
    "in", "include", "individual", "inout", "interface", "istype", "item",
    "join", "language", "library", "locale", "loop", "merge", "message",
    "meta", "metadata", "nonunique", "not", "null", "objective", "occurrence",
    "of", "or", "ordered", "out", "package", "parallel", "part", "perform",
    "port", "private", "protected", "public", "readonly", "redefines", "ref",
    "references", "rendering", "rep", "require", "requirement", "return",
    "satisfy", "send", "snapshot", "specializes", "stakeholder", "standard",
    "state", "subject", "subsets", "succession", "then", "timeslice", "to",
    "transition", "until", "use", "variant", "variation", "verification",
    "verify", "via", "view", "viewpoint", "when", "while", "xor",
}
IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


# ------------------------------------------------------------------- helpers

def unfold(s):
    """TSV cells encode real line breaks as the literal token <NL>."""
    return (s or "").replace("<NL>", "\n").strip()


def q(name):
    """Quote a SysML v2 name if needed (backslash escaping, as used by the
    CATIA 2026x textual export and the Pilot Implementation)."""
    name = name if name is not None else ""
    if IDENT_RE.match(name) and name not in RESERVED:
        return name
    body = (name.replace("\\", "\\\\").replace("'", "\\'")
                .replace("\n", "\\n").replace("\t", "\\t").replace("\r", ""))
    return f"'{body}'"


def qpath(qn):
    """Render a qualified name as a v2 path with per-segment quoting."""
    segs = [s for s in unfold(qn).split("::") if s.strip()]
    return "::".join(q(s) for s in segs)


def last_seg(qn):
    segs = [s for s in unfold(qn).split("::") if s.strip()]
    return segs[-1] if segs else ""


def mult_str(m):
    m = (m or "").strip()
    if not m or ".." not in m:
        return ""
    lo, hi = m.split("..", 1)
    if hi in ("-1", "*"):
        hi = "*"
    if lo == hi:
        return "" if lo == "1" else f"[{lo}]"
    return f"[{lo}..{hi}]"


def base_type(et):
    """'Block [Class]' -> 'Block';  'ItemFlow [InformationFlow]' -> 'ItemFlow'"""
    return re.sub(r"\s*\[.*\]\s*$", "", et or "").strip()


def pretty_metatype(et):
    """Report-style metatype: CATIA uses spaced names ('Activity Parameter
    Node'); split camel-case when the label has no space already."""
    b = base_type(et)
    if " " in b:
        return b
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", b)


class Node:
    __slots__ = ("row", "et", "name", "qn", "owner", "eid", "children",
                 "status", "dropped", "extra")

    def __init__(self, row, idx):
        def g(col):
            j = idx.get(col)
            return row[j].strip() if j is not None and j < len(row) else ""
        self.row = {c: g(c) for c in idx}
        self.et = base_type(g(HDR_TYPE))
        self.name = unfold(g("Name"))
        self.qn = g("Qualified Name")
        self.owner = g("Owner Qualified Name")
        self.eid = g("Element ID")
        self.children = []
        self.status = "pending"   # migrated | consumed | notmig | diagram
        self.dropped = []         # property names conveyed only as comments/not at all
        self.extra = {}

    def g(self, col):
        return self.row.get(col, "")


# ----------------------------------------------------------------- TSV load

def load_tsv(path):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=enc) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        sys.exit(f"Cannot decode {path}")
    rows = [ln.split("\t") for ln in text.splitlines() if ln.strip()]
    hdr = next((r for r in rows if HDR_TYPE in r), None)
    if hdr is None:
        sys.exit(f"{path}: not a SysML17_Migration_Export TSV (header not found)")
    idx = {h.strip(): i for i, h in enumerate(hdr)}
    nodes, dropped = [], 0
    for r in rows[rows.index(hdr) + 1:]:
        if len(r) < 6 or (not r[1].strip() and not r[5].strip()):
            dropped += 1
            continue
        nodes.append(Node(r, idx))
    if dropped:
        print(f"note: skipped {dropped} blank/fragment TSV rows")
    return nodes


def build_tree(nodes):
    """Order-independent containment rebuild: attach every node to a node
    whose Qualified Name equals its Owner Qualified Name. Works whether the
    TSV is still in the template's depth-first order or has been re-sorted
    (e.g. by element type in Excel). When several nodes share a QN (unnamed
    elements), the nearest preceding one is preferred, which reproduces the
    depth-first behaviour on unsorted files."""
    by_qn = {}
    for idx, n in enumerate(nodes):
        n.extra["row"] = idx
        by_qn.setdefault(n.qn, []).append(n)

    roots = []
    for n in nodes:
        if n.et == "Model":
            n.status = "consumed"
            continue
        # diagrams and the Model row can never own model elements - a diagram
        # merely named like a package must not capture that package's children
        cands = [c for c in by_qn.get(n.owner, [])
                 if c is not n and c.et not in ("Model", "Diagram")]
        parent = None
        if cands:
            before = [c for c in cands if c.extra["row"] < n.extra["row"]]
            parent = before[-1] if before else cands[0]
        if parent is not None:
            parent.children.append(n)
            n.extra["parent"] = parent
        else:
            roots.append(n)

    # safety net: a QN collision could create an ownership cycle that leaves
    # a subtree unreachable - promote such nodes to roots (each node must be
    # emitted exactly once)
    reach = set()

    def mark(m):
        if id(m) in reach:
            return
        reach.add(id(m))
        for c in m.children:
            mark(c)

    for r in roots:
        mark(r)
    for n in nodes:
        if n.et != "Model" and id(n) not in reach:
            p = n.extra.get("parent")
            if p is not None and n in p.children:
                p.children.remove(n)
            roots.append(n)
            mark(n)
    return roots


# ------------------------------------------------------------------ emitter

class Emitter:
    # emitted after their siblings so referenced features exist first
    EDGE_ETS = ("ControlFlow", "ObjectFlow", "Transition",
                "Connector", "BindingConnector", "ItemFlow",
                "DeriveReqt", "Satisfy")

    # 1.7 element types that become v2 DEFINITIONS (not usages/features)
    DEF_ETS = {"Package", "Profile", "Block", "Class", "Actor",
               "InterfaceBlock", "ConstraintBlock", "ValueType", "DataType",
               "Enumeration", "Signal", "UseCase", "TestCase",
               "StateMachine", "Activity"}

    def __init__(self, dialect, model_roots, et_by_qn=None, node_by_qn=None):
        self.d = dialect
        self.lines = []
        self.roots = {last_seg(r.qn) for r in model_roots}
        self.anon = 0
        self.scope = []   # stack of per-container declared usage-name sets
        self.path = []    # names of the containers we are emitting inside
        self.et_by_qn = et_by_qn or {}
        self.node_by_qn = node_by_qn or {}
        self.nested = {}  # owner usage name -> names of its nested features
                          # (pins/payloads), for succession dot-path endpoints

    def declare(self, name):
        if self.scope and name:
            self.scope[-1].add(name)

    def resolves(self, qn):
        return bool(self.scope) and last_seg(qn) in self.scope[-1]

    def relref(self, qn):
        """Reference text for a model element: strip the segments shared with
        the current emission scope. Short references resolve through the
        enclosing packages and dodge name shadowing (e.g. a part def named
        like its owning package breaks absolute 'Root::...' paths)."""
        segs = [s for s in unfold(qn).split("::") if s.strip()]
        if not segs:
            return ""
        cur = [s for s in self.path if s]
        k = 0
        while k < min(len(segs) - 1, len(cur)) and segs[k] == cur[k]:
            k += 1
        return "::".join(q(s) for s in segs[k:])

    def w(self, depth, text):
        self.lines.append("    " * depth + text)

    def typeref(self, tqn, node, prop="type"):
        """Return (': path', comment) for a type reference; external types
        that cannot resolve become a TODO comment instead."""
        tqn = unfold(tqn)
        if not tqn:
            return "", ""
        root = tqn.split("::")[0].strip()
        if root in self.roots:
            return f" : {self.relref(tqn)}", ""
        prim = PRIMITIVES.get(last_seg(tqn))
        if prim:
            return f" : {prim}", ""
        node.dropped.append((prop, f"type '{tqn}' is outside the exported "
                             "model and not a known ScalarValues primitive - "
                             "re-type manually"))
        return "", f" // TODO-MIGRATE {prop} was {tqn}"

    # -- endpoint dot-path for connectors: part.role relative to owner block
    def endpath(self, role_qn, pwp_qn):
        role = last_seg(role_qn)
        pwp = last_seg(pwp_qn)
        if not role:
            return ""
        return f"{q(pwp)}.{q(role)}" if pwp else q(role)

    def todo(self, depth, node, why):
        nm = f" {q(node.name)}" if node.name else ""
        self.w(depth, f"// TODO-MIGRATE {pretty_metatype(node.et)}{nm}: {why}")
        node.status = "notmig"
        node.extra["reason"] = why

    # ------------------------------------------------------------------ visit
    def visit(self, n, depth):
        et = n.et
        han = {
            "Package": self.pkg, "Profile": self.pkg, "Model": self.pkg,
            "Block": self.block, "Class": self.block, "Actor": self.block,
            "InterfaceBlock": self.portdef, "ConstraintBlock": self.constraintdef,
            "ValueType": self.attrdef, "DataType": self.attrdef,
            "Enumeration": self.enumdef, "EnumerationLiteral": self.enumlit,
            "Signal": self.itemdef, "Requirement": self.requirement,
            "UseCase": self.usecase, "TestCase": self.verification,
            "PartProperty": self.part, "ReferenceProperty": self.refpart,
            "ValueProperty": self.attr, "FlowProperty": self.flowprop,
            "ConstraintProperty": self.assertconstraint,
            "ConstraintParameter": self.attr, "Property": self.plainprop,
            "ProxyPort": self.port, "Port": self.port, "FlowPort": self.port,
            "Connector": self.connector, "BindingConnector": self.binding,
            "ConnectorEnd": self.consumed, "Association": self.association,
            "ItemFlow": self.itemflow, "Generalization": self.consumed,
            "StateMachine": self.statemachine, "Region": self.region,
            "State": self.state, "Final State": self.finalstate,
            "Pseudostate": self.pseudostate, "Transition": self.transition,
            "Trigger": self.absorbed, "SignalEvent": self.absorbed,
            "Activity": self.activity, "Parameter": self.parameter,
            "ActivityParameterNode": self.paramnode,
            "CallBehaviorAction": self.cba,
            "InputPin": self.pin, "Input Pin": self.pin,
            "OutputPin": self.pin, "Output Pin": self.pin,
            "ControlFlow": self.edge, "ObjectFlow": self.edge,
            "Initial Node": self.initialnode, "InitialNode": self.initialnode,
            "Activity Final Node": self.finalnode,
            "ActivityFinalNode": self.finalnode,
            "Merge Node": self.controlnode, "MergeNode": self.controlnode,
            "Decision Node": self.controlnode,
            "DecisionNode": self.controlnode,
            "Fork Node": self.controlnode, "ForkNode": self.controlnode,
            "Join Node": self.controlnode, "JoinNode": self.controlnode,
            "AcceptEventAction": self.acceptaction,
            "Accept Event Action": self.acceptaction,
            "SendSignalAction": self.sendaction,
            "Send Signal Action": self.sendaction,
            "Satisfy": self.satisfy, "DeriveReqt": self.dependencyrel,
            "Abstraction": self.dependencyrel, "Dependency": self.dependencyrel,
            "Usage": self.dependencyrel,
            "Diagram": self.diagram,
        }.get(et, self.unsupported)
        han(n, depth)

    def children(self, n, depth):
        for c in n.children:
            self.visit(c, depth)

    def block_body(self, n, depth, headline, comment=""):
        """Emit `headline { children }` or `headline;`"""
        gens = [c for c in n.children if c.et == "Generalization"]
        sup = ""
        for gz in gens:
            gq = gz.g("General/Parent QN")
            if gq:
                sup += f" :> {qpath(gq)}"
            gz.status = "consumed"
        kids = [c for c in n.children if c.et != "Generalization"]
        if kids:
            self.w(depth, f"{headline}{sup} {{{comment}")
            # declare-before-use: emit edges/connectors after everything else
            # so their endpoint names can be resolved against declared usages
            self.scope.append(set())
            self.path.append(n.name)
            plain = [c for c in kids if c.et not in self.EDGE_ETS]
            edges = [c for c in kids if c.et in self.EDGE_ETS]
            for c in plain:
                self.visit(c, depth + 1)
            for c in edges:
                self.visit(c, depth + 1)
            self.path.pop()
            self.scope.pop()
            self.w(depth, "}")
        else:
            self.w(depth, f"{headline}{sup};{comment}")
        n.status = "migrated"

    # ---------------------------------------------------------- structural
    def pkg(self, n, depth):
        note = ""
        if n.et == "Profile":
            note = " // TODO-MIGRATE was a UML Profile (stereotypes have no v2 home)"
        self.block_body(n, depth, f"package {q(n.name)}", note)

    def block(self, n, depth):
        note = " // actor" if n.et == "Actor" else (
            " // was plain UML Class" if n.et == "Class" else "")
        self.block_body(n, depth, f"part def {q(n.name)}", note)

    def portdef(self, n, depth):
        self.block_body(n, depth, f"port def {q(n.name)}")

    def constraintdef(self, n, depth):
        self.block_body(n, depth, f"constraint def {q(n.name)}")

    def attrdef(self, n, depth):
        self.block_body(n, depth, f"attribute def {q(n.name)}")

    def enumdef(self, n, depth):
        self.block_body(n, depth, f"enum def {q(n.name)}")

    def enumlit(self, n, depth):
        self.w(depth, f"{q(n.name)};")
        n.status = "migrated"

    def itemdef(self, n, depth):
        self.block_body(n, depth, f"item def {q(n.name)}")

    def usecase(self, n, depth):
        self.block_body(n, depth, f"use case def {q(n.name)}")

    def verification(self, n, depth):
        self.block_body(n, depth, f"verification def {q(n.name)}",
                        " // was SysML 1.7 TestCase")

    def requirement(self, n, depth):
        self.w(depth, f"requirement {q(n.name)} {{")
        self.w(depth + 1, "doc /* TODO-MIGRATE requirement text not exported (GAP-05) */")
        n.dropped.append(("text", "requirement text/Id tag are not exported "
                          "by the template (GAP-05) - re-enter or import from "
                          "the original source"))
        self.path.append(n.name)
        self.children(n, depth + 1)
        self.path.pop()
        self.w(depth, "}")
        n.status = "migrated"

    # ------------------------------------------------------------- features
    def part(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        m = mult_str(n.g("Multiplicity"))
        m = f" {m}" if m else ""
        self.declare(n.name)
        self.w(depth, f"part {q(n.name)}{m}{t};{cm}")
        n.status = "migrated"

    def refpart(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        m = mult_str(n.g("Multiplicity"))
        m = f" {m}" if m else ""
        self.declare(n.name)
        self.w(depth, f"ref part {q(n.name)}{m}{t};{cm}")
        n.status = "migrated"

    def attr(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        m = mult_str(n.g("Multiplicity"))
        m = f" {m}" if m else ""
        self.declare(n.name)
        self.w(depth, f"attribute {q(n.name)}{m}{t};{cm}")
        n.status = "migrated"

    def flowprop(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        d = n.g("Direction").lower()
        d = d if d in ("in", "out", "inout") else "inout"
        self.declare(n.name)
        self.w(depth, f"{d} item {q(n.name)}{t};{cm}")
        n.status = "migrated"

    def assertconstraint(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        self.declare(n.name)
        self.w(depth, f"assert constraint {q(n.name)}{t};{cm}")
        n.status = "migrated"

    def port(self, n, depth):
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        self.declare(n.name)
        self.w(depth, f"port {q(n.name)}{t};{cm}")
        if n.g("Direction"):
            n.dropped.append(("direction", "no direct v2 syntax for a "
                              "directed port usage - review conjugation (~) "
                              "manually"))
        n.status = "migrated"

    def plainprop(self, n, depth):
        # association-owned member ends are covered by the Association row
        if not n.name and not n.g("Type Qualified Name"):
            n.status = "consumed"
            return
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        m = mult_str(n.g("Multiplicity"))
        m = f" {m}" if m else ""
        kw = "ref part" if t and " : " in t and not t.endswith(
            ("Real", "Integer", "String", "Boolean")) else "attribute"
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        self.w(depth, f"{kw} {q(raw)}{m}{t};{cm}")
        n.status = "migrated"

    def anon_name(self):
        self.anon += 1
        return f"__unnamed{self.anon}"

    # ----------------------------------------------------------- connectors
    def connector(self, n, depth):
        a = self.endpath(n.g("CE0 Role QN"), n.g("Src End Part With Port QN"))
        b = self.endpath(n.g("CE1 Role QN"), n.g("Tgt End Part With Port QN"))
        if a and b:
            nm = f"connection {q(n.name)} " if n.name else "connection "
            self.w(depth, f"{nm}connect {a} to {b};")
            n.status = "migrated"
        else:
            self.todo(depth, n, "connector endpoints unresolved in TSV")
        for c in n.children:
            c.status = "consumed"

    def binding(self, n, depth):
        a = self.endpath(n.g("CE0 Role QN"), n.g("Src End Part With Port QN"))
        b = self.endpath(n.g("CE1 Role QN"), n.g("Tgt End Part With Port QN"))
        if a and b:
            head = f"binding {q(n.name)} bind " if n.name else "bind "
            self.w(depth, f"{head}{a} = {b};")
            n.status = "migrated"
        else:
            self.todo(depth, n, "binding endpoints unresolved in TSV")
        for c in n.children:
            c.status = "consumed"

    def association(self, n, depth):
        nm = q(n.name) if n.name else self.anon_name()
        ends = []
        for col in ("Association End 0", "Association End 1"):
            raw = n.g(col)
            if not raw:
                continue
            bits = [b.strip() for b in raw.split("::")]
            # role :: typeQN(may itself contain ::) :: mult :: aggregation
            role = bits[0] if bits else ""
            agg = bits[-1] if len(bits) >= 4 else ""
            mult = bits[-2] if len(bits) >= 4 else ""
            tqn = "::".join(bits[1:-2]) if len(bits) >= 4 else ""
            ends.append((role, tqn, mult, agg))
        if ends:
            self.w(depth, f"connection def {nm} {{")
            for i, (role, tqn, mult, agg) in enumerate(ends):
                t, cm = self.typeref(tqn, n, prop=f"end{i} type")
                m = mult_str(mult)
                m = f" {m}" if m else ""
                rn = q(role) if role else f"end{i}"
                note = f" // aggregation: {agg}" if agg and agg != "none" else ""
                self.w(depth + 1, f"end {rn}{m}{t};{cm}{note}")
                if agg and agg != "none":
                    n.dropped.append((f"end{i} aggregation",
                                      "v2 connection ends do not encode the "
                                      "aggregation kind - conveyed as a "
                                      "comment only"))
            self.w(depth, "}")
            n.status = "migrated"
        else:
            self.todo(depth, n, "association ends missing in TSV")
        for c in n.children:
            if c.et == "Property":
                c.status = "consumed"
            else:
                self.visit(c, depth)

    def itemflow(self, n, depth):
        conv = n.g("Item Conveyed Qualified Name")
        src, tgt = n.g("Source/Client Qualified Name"), n.g("Target/Supplier Qualified Name")
        if self.d == "msosa" and src and tgt:
            of = f" of {self.relref(conv)}" if conv else ""
            self.w(depth, f"flow{of} from {self.relref(src)} to {self.relref(tgt)};")
            n.status = "migrated"
        else:
            # pilot needs port-to-port dot-paths (realizing connector not in TSV)
            self.todo(depth, n,
                      f"ItemFlow conveys {conv or '?'} from {src or '?'} to "
                      f"{tgt or '?'} - re-express as port dot-path flow (FIX-6)")
            why = ("pilot flows require port-to-port dot-paths; the realizing "
                   "connector is not exported in the TSV (FIX-6)")
            n.dropped += [("conveyed", why), ("informationSource", why),
                          ("informationTarget", why)]

    # -------------------------------------------------------------- states
    def statemachine(self, n, depth):
        self.block_body(n, depth, f"state def {q(n.name)}")

    def region(self, n, depth):
        plain = [c for c in n.children if c.et not in self.EDGE_ETS]
        edges = [c for c in n.children if c.et in self.EDGE_ETS]
        if n.name:
            # a NAMED region keeps its own level as a nested composite state
            # so v1.7 qualified names (…::region::state) survive migration
            self.declare(n.name)
            self.w(depth, f"state {q(n.name)} {{ // was SysML 1.7 Region")
            self.scope.append(set())
            self.path.append(n.name)
            for c in plain:
                self.visit(c, depth + 1)
            for c in edges:
                self.visit(c, depth + 1)
            self.path.pop()
            self.scope.pop()
            self.w(depth, "}")
            n.status = "migrated"
        else:
            # unnamed: transparent, keep declare-before-use ordering
            n.status = "consumed"
            for c in plain:
                self.visit(c, depth)
            for c in edges:
                self.visit(c, depth)

    def state(self, n, depth):
        self.declare(n.name)
        head = f"state {q(n.name)}"
        body = []
        for col, kw in (("State Entry Behavior QN", "entry"),
                        ("State Do Activity QN", "do"),
                        ("State Exit Behavior QN", "exit")):
            bqn = n.g(col)
            if bqn:
                if self.d == "msosa":
                    body.append(f"{kw} action : {self.relref(bqn)};")
                else:
                    body.append(f"{kw} action on{kw.capitalize()}; "
                                f"// performs {bqn}")
                    n.dropped.append((f"{kw} behavior binding",
                                      "pilot dialect requires a named "
                                      "placeholder action; the performed "
                                      "behavior is referenced in a trailing "
                                      "comment - bind manually"))
        kids = list(n.children)
        if body or kids:
            self.w(depth, head + " {")
            for b in body:
                self.w(depth + 1, b)
            self.scope.append(set())
            self.path.append(n.name)
            plain = [c for c in kids if c.et not in self.EDGE_ETS]
            edges = [c for c in kids if c.et in self.EDGE_ETS]
            for c in plain:
                self.visit(c, depth + 1)
            for c in edges:
                self.visit(c, depth + 1)
            self.path.pop()
            self.scope.pop()
            self.w(depth, "}")
        else:
            self.w(depth, head + ";")
        n.status = "migrated"

    def finalstate(self, n, depth):
        nm = n.name if n.name else "done"
        self.declare(nm)
        self.w(depth, f"state {q(nm)}; // final state")
        n.status = "migrated"

    def pseudostate(self, n, depth):
        kind = n.g("Pseudostate Kind") or "initial"
        if kind == "initial":
            if n.name:
                # keep the v1.7 name: the initial pseudostate becomes a plain
                # state routed from 'entry'; the initial transition out of it
                # keeps its own name too (see transition())
                self.declare(n.name)
                self.w(depth, f"state {q(n.name)}; // was initial pseudostate")
                n.status = "migrated"
            else:
                n.status = "consumed"   # rendered by the initial transition
        else:
            self.todo(depth, n, f"pseudostate kind '{kind}' needs manual mapping")

    def transition(self, n, depth):
        src, tgt = n.g("Source Node Qualified Name"), n.g("Target Node Qualified Name")
        eff = n.g("Transition Effect QN")

        # A v2 transition is a one-line usage and cannot own definitions.
        # Hoist non-trigger children (e.g. an effect Activity owned by the
        # transition, incl. internal transitions typed by an activity) into
        # the enclosing scope FIRST so their whole subtree migrates; only
        # trigger machinery stays consumed. Note the hoisted definition's
        # qualified name shifts one level up (transition level dropped).
        eff_name = ""
        for c in n.children:
            if c.et in ("Trigger", "SignalEvent"):
                c.status = "consumed"
                continue
            self.visit(c, depth)
            if c.name and unfold(c.qn) == unfold(eff):
                eff_name = c.name

        if not (src and tgt):
            self.todo(depth, n, "transition endpoints missing")
            return
        trig, grd = n.g("Trigger"), n.g("Guard Expression")
        cm = ""
        notes = []
        if trig:
            notes.append(f"trigger: {trig}")
            n.dropped.append(("trigger", "accept clause not auto-emitted "
                              "(conservative) - conveyed as comment; re-add "
                              "as 'accept <signal>' manually"))
        if grd:
            notes.append(f"guard: [{grd}]")
            n.dropped.append(("guard", "guard expression conveyed as comment "
                              "- re-add as 'if <expr>' manually"))
        do_txt = ""
        if eff:
            if eff_name:
                # the owned effect was hoisted just above as an action def -
                # reference it from the transition's do clause
                do_txt = f"do action {q(eff_name + '_do')} : {q(eff_name)} "
                n.dropped.append(("effect owner", "effect Activity was owned "
                                  "by the transition and was hoisted into "
                                  "the enclosing scope as an action def (v2 "
                                  "transitions cannot own definitions) - its "
                                  "qualified name shifted one level up"))
            else:
                notes.append(f"effect: {eff}")
                n.dropped.append(("effect", "effect behavior conveyed as "
                                  "comment - re-add as a 'do action' manually"))
        if notes:
            cm = " // TODO-MIGRATE " + "; ".join(notes)
        # initial transition?
        srcnode_kind = n.extra.get("src_kind", "")
        if srcnode_kind == "initial":
            if self.resolves(src):
                # the initial pseudostate was declared as a state: enter
                # through it, then keep the transition with its v1.7 name
                if not self.resolves(tgt):
                    self.todo(depth, n, f"initial transition target "
                                        f"'{last_seg(tgt)}' not a declared state")
                    return
                self.w(depth, f"entry; then {q(last_seg(src))};")
                nm = f" {q(n.name)}" if n.name else ""
                self.w(depth, f"transition{nm} first {q(last_seg(src))} "
                              f"{do_txt}then {q(last_seg(tgt))};{cm}")
            elif not self.resolves(tgt):
                self.todo(depth, n, f"initial transition target '{last_seg(tgt)}' "
                                    "not a declared state")
                return
            else:
                if do_txt:
                    cm += f" // effect hoisted as action def {q(eff_name)}"
                self.w(depth, f"entry; then {q(last_seg(tgt))};{cm}")
        else:
            if not (self.resolves(src) and self.resolves(tgt)):
                self.todo(depth, n, f"endpoints '{last_seg(src)}' -> "
                                    f"'{last_seg(tgt)}' not declared usages "
                                    "(unnamed/pseudo node?)")
                return
            nm = f" {q(n.name)}" if n.name else ""
            self.w(depth, f"transition{nm} first {q(last_seg(src))} "
                          f"{do_txt}then {q(last_seg(tgt))};{cm}")
        n.status = "migrated"

    # ------------------------------------------------------------ behavior
    def activity(self, n, depth):
        self.block_body(n, depth, f"action def {q(n.name)}")

    def parameter(self, n, depth):
        key = (n.owner, n.name)
        if key in getattr(self, "_seen_params", set()):
            n.status = "consumed"
            return
        self._seen_params = getattr(self, "_seen_params", set())
        self._seen_params.add(key)
        d = (n.g("Direction") or "in").lower()
        d = d if d in ("in", "out", "inout") else "in"
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        self.declare(n.name)
        self.w(depth, f"{d} {q(n.name)}{t};{cm}")
        n.status = "migrated"

    def paramnode(self, n, depth):
        self.parameter(n, depth)

    def cba(self, n, depth):
        called = n.g("Called Behavior Qualified Name")
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        head = f"action {q(raw)}"
        if called:
            head += f" : {self.relref(called)}"
        if n.children:
            self.w(depth, head + " {")
            self.scope.append(set())
            self.path.append(raw)
            for c in n.children:
                self.visit(c, depth + 1)
            self.path.pop()
            self.scope.pop()
            self.w(depth, "}")
        else:
            self.w(depth, head + ";")
        n.status = "migrated"

    def pin(self, n, depth):
        d = "out" if n.et in ("OutputPin", "Output Pin") else "in"
        raw = n.name if n.name else self.anon_name()
        t, cm = self.typeref(n.g("Type Qualified Name"), n)
        self.declare(raw)
        owner = self.path[-1] if self.path else ""
        if owner:
            self.nested.setdefault(owner, set()).add(raw)
        self.w(depth, f"{d} {q(raw)}{t};{cm}")
        n.status = "migrated"

    def initialnode(self, n, depth):
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        self.w(depth, f"action {q(raw)}; // was SysML 1.7 Initial Node")
        self.w(depth, f"first start then {q(raw)};")
        n.status = "migrated"

    def finalnode(self, n, depth):
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        self.w(depth, f"action {q(raw)}; // was SysML 1.7 Activity Final Node")
        self.w(depth, f"first {q(raw)} then done;")
        n.status = "migrated"

    CONTROL_KW = {"Merge Node": "merge", "MergeNode": "merge",
                  "Decision Node": "decide", "DecisionNode": "decide",
                  "Fork Node": "fork", "ForkNode": "fork",
                  "Join Node": "join", "JoinNode": "join"}

    def controlnode(self, n, depth):
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        self.w(depth, f"{self.CONTROL_KW[n.et]} {q(raw)};")
        n.status = "migrated"
        for c in n.children:
            self.visit(c, depth)

    def acceptaction(self, n, depth):
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        trig = next((c for c in n.children if c.et == "Trigger"), None)
        outpin = next((c for c in n.children
                       if c.et in ("OutputPin", "Output Pin")), None)
        sig_qn = ""
        if trig is not None:
            ev = self.node_by_qn.get(trig.g("Signal Qualified Name"))
            if ev is not None:
                sig_qn = ev.g("Signal Qualified Name")
                ev.extra["absorbed_by"] = raw
            trig.status = "notmig"
            trig.extra["reason"] = (f"absorbed into accept action {q(raw)} - "
                                    "no separate v2 element")
        payload = outpin.name if (outpin is not None and outpin.name) else "evt"
        if sig_qn:
            t, cm = self.typeref(sig_qn, n, prop="accepted signal")
            # head form: the payload becomes a parameter of the accept action
            # itself, so succession-flow dot-paths ('57'.result) resolve
            self.w(depth, f"action {q(raw)} accept {q(payload)}{t};{cm}")
            if outpin is not None:
                outpin.status = "consumed"   # matched as the accept payload
            self.nested.setdefault(raw, set()).add(payload)
        else:
            self.w(depth, f"action {q(raw)}; // TODO-MIGRATE accept event: "
                          "trigger/signal chain not resolvable from TSV")
            n.dropped.append(("accept event", "trigger/signal chain not "
                              "resolvable - re-add 'accept <payload> : "
                              "<Signal>' manually"))
        for c in n.children:
            if c.status == "pending":
                self.visit(c, depth)
        n.status = "migrated"

    def sendaction(self, n, depth):
        raw = n.name if n.name else self.anon_name()
        self.declare(raw)
        sig = n.g("Signal Qualified Name")
        was = f" of {last_seg(sig)}" if sig else ""
        self.w(depth, f"action {q(raw)}; // TODO-MIGRATE send{was}: signal/"
                      "target not exported by the template - re-add "
                      "'send new <Signal>() to <target>' manually")
        n.dropped.append(("sent signal", "signal and send target are not "
                          "exported by the template - re-add the send "
                          "clause manually"))
        n.status = "migrated"
        for c in n.children:
            if c.status == "pending":
                self.visit(c, depth)

    def absorbed(self, n, depth):
        """1.7 helper elements (triggers, signal events) whose content is
        carried by an enclosing v2 construct: no separate v2 element exists.
        Named ones get a report row so nothing is silently dropped."""
        if n.name:
            n.status = "notmig"
            if "reason" not in n.extra:
                by = n.extra.get("absorbed_by")
                n.extra["reason"] = (
                    (f"content absorbed by v2 construct {q(by)}" if by
                     else "content absorbed by the enclosing v2 construct")
                    + " - no separate named v2 element")
        else:
            n.status = "consumed"

    def _endref(self, qn):
        """Succession endpoint: a declared sibling usage, or an owner.member
        dot-path for pins/payloads declared one level down."""
        if not qn:
            return ""
        last = last_seg(qn)
        if self.resolves(qn):
            return q(last)
        segs = [s for s in unfold(qn).split("::") if s.strip()]
        if len(segs) >= 2:
            owner = segs[-2]
            if (self.scope and owner in self.scope[-1]
                    and last in self.nested.get(owner, ())):
                return f"{q(owner)}.{q(last)}"
        return ""

    def edge(self, n, depth):
        src, tgt = n.g("Source Node Qualified Name"), n.g("Target Node Qualified Name")
        a, b = self._endref(src), self._endref(tgt)
        if a and b:
            grd = n.g("Guard Expression")
            cm = ""
            if grd:
                cm = f" // TODO-MIGRATE guard: [{grd}]"
                n.dropped.append(("guard", "guard expression conveyed as "
                                  "comment - re-add manually"))
            nm = f" {q(n.name)}" if n.name else ""
            if n.et == "ObjectFlow":
                self.w(depth, f"succession flow{nm} from {a} to {b};{cm}")
            elif nm:
                self.w(depth, f"succession{nm} first {a} then {b};{cm}")
            else:
                self.w(depth, f"first {a} then {b};{cm}")
            n.status = "migrated"
        else:
            self.todo(depth, n, f"edge '{last_seg(src) or '?'}' -> "
                                f"'{last_seg(tgt) or '?'}' references "
                                "undeclared/unnamed nodes (pin, initial, final...)")

    # -------------------------------------------------------- traceability
    def dependencyrel(self, n, depth):
        """SysML 1.7 Abstraction / Dependency / Usage -> v2 `dependency`
        (a first-class KerML construct: dependency [name] from a to b;)."""
        src = n.g("Source/Client Qualified Name")
        tgt = n.g("Target/Supplier Qualified Name")
        if src and tgt:
            srcs = ", ".join(self.relref(s) for s in src.split(", ") if s.strip())
            tgts = ", ".join(self.relref(t) for t in tgt.split(", ") if t.strip())
            nm = f" {q(n.name)}" if n.name else ""
            cm = ("" if n.et == "Dependency"
                  else f" // was SysML 1.7 {pretty_metatype(n.et)}")
            self.w(depth, f"dependency{nm} from {srcs} to {tgts};{cm}")
            n.status = "migrated"
        else:
            self.todo(depth, n, "client/supplier missing in TSV")

    def satisfy(self, n, depth):
        req = n.g("Target/Supplier Qualified Name")
        by = n.g("Source/Client Qualified Name")
        if req and by:
            by_et = self.et_by_qn.get(by, "")
            if self.d == "pilot" and by_et in self.DEF_ETS:
                # v2 'satisfy' needs a usage target; the 1.7 client is a
                # definition - keep the element (and its name) as a plain
                # dependency instead of dropping it
                nm = f" {q(n.name)}" if n.name else ""
                self.w(depth, f"dependency{nm} from {self.relref(by)} to "
                              f"{self.relref(req)}; // was SysML 1.7 Satisfy "
                              "(v2 satisfy needs a usage target)")
                n.dropped.append(("satisfy semantics", "emitted as a plain "
                                  "dependency: v2 'satisfy' requires a usage "
                                  f"target but the 1.7 client is a {by_et} "
                                  "definition - re-express with 'satisfy ... "
                                  "by <part usage>' once a usage exists"))
                n.status = "migrated"
                return
            self.w(depth, f"satisfy {self.relref(req)} by {self.relref(by)};")
            n.status = "migrated"
        else:
            self.todo(depth, n, "satisfy client/supplier missing")

    # ------------------------------------------------------------- residue
    def diagram(self, n, depth):
        dt = n.g("Diagram Type") or "Diagram"
        self.w(depth, f"// DIAGRAM (not migrated): {n.name}  [{dt}]")
        n.status = "diagram"
        # diagrams own nothing in a healthy export, but never drop a subtree
        for c in n.children:
            self.visit(c, depth)

    def consumed(self, n, depth):
        n.status = "consumed"

    def unsupported(self, n, depth):
        why = {
            "Interface": "UML Interface - map manually (port def / interface def)",
            "Flow Specification": "FlowSpecification - superseded; port def with items",
            "Interface Realization": "no direct v2 construct",
            "Operation": "operations have no direct v2 home (GAP-14)",
            "Interaction": "sequence diagrams not convertible from element data",
            "Message": "sequence message - not convertible",
            "Lifeline": "sequence lifeline - not convertible",
            "Allocate": "v2 allocates usages, 1.7 targeted defs - decide targets",
            "Verify": "belongs in a verification def (FIX-3)",
            "Constraint": "opaque constraint - re-express as require/assert",
            "Activity Partition": "swim lanes have no v2 construct - carry "
                                  "the intent with allocate/perform instead",
            "Central Buffer Node": "no direct v2 construct - consider an "
                                   "item/store pattern",
            "Legend": "diagram legend artifact, not model semantics",
            "Legend Item": "diagram legend artifact, not model semantics",
            "Instance Value": "instance specifications not exported in full",
            "Opaque Behavior": "body text not exported",
        }.get(n.et, "no automatic v2 mapping")
        self.todo(depth, n, why)
        for c in n.children:
            self.visit(c, depth)


def mark_initial_transitions(nodes):
    kind_by_qn = {}
    for n in nodes:
        if n.et == "Pseudostate":
            kind_by_qn[n.qn] = n.g("Pseudostate Kind") or "initial"
    for n in nodes:
        if n.et == "Transition":
            n.extra["src_kind"] = kind_by_qn.get(
                n.g("Source Node Qualified Name"), "")


# ------------------------------------------------------------------- report

def write_report(nodes, out_xlsx, project):
    import openpyxl
    from openpyxl.styles import Font

    def url(n):
        qs = urllib.parse.urlencode({
            "projectName": project,
            "elementName": f"{pretty_metatype(n.et)} {n.name}".strip(),
            "metaType": n.et,
            "elementQName": unfold(n.qn)})
        return f"mdel://{n.eid}?{qs}" if n.eid else ""

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Not Migrated Elements"
    ws.append(["Element MetaType", "Element Name", "Qualified Name",
               "Element Url", "Element ID", "Reason"])
    for n in nodes:
        if n.status == "notmig":
            ws.append([pretty_metatype(n.et), n.name, unfold(n.qn),
                       url(n), n.eid,
                       n.extra.get("reason", "no automatic v2 mapping")])

    ws2 = wb.create_sheet("Not Migrated Properties")
    ws2.append(["Element Name", "Element MetaType", "Property Name",
                "Qualified Name", "Element Url", "Reason"])
    for n in nodes:
        if n.status == "migrated" and n.dropped:
            first = True
            for entry in dict.fromkeys(n.dropped):
                prop, why = entry if isinstance(entry, tuple) else (entry, "")
                if first:
                    ws2.append([n.name, pretty_metatype(n.et), prop,
                                unfold(n.qn), url(n), why])
                    first = False
                else:
                    ws2.append([None, None, prop, None, None, why])

    DIAGRAM_REASON = ("SysML v2 textual notation has no diagram persistence; "
                      "layout/x-y positions are not exported (GAP-02) - "
                      "regenerate the diagram in the target tool from the "
                      "migrated elements")
    ws3 = wb.create_sheet("Not Migrated Diagrams")
    ws3.append(["Diagram Type", "Element Name", "Qualified Name",
                "Element Url", "Element ID", "Reason"])
    for n in nodes:
        if n.status == "diagram":
            ws3.append([n.g("Diagram Type") or "Diagram", n.name,
                       unfold(n.qn), url(n), n.eid, DIAGRAM_REASON])

    for ws_ in (ws, ws2, ws3):
        for c in ws_[1]:
            c.font = Font(name="Arial", size=10, bold=True)
        for row in ws_.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
        ws_.freeze_panes = "A2"
        for col in ws_.columns:
            width = max((len(str(c.value)) for c in col if c.value), default=10)
            ws_.column_dimensions[col[0].column_letter].width = min(width + 2, 80)
    wb.save(out_xlsx)


# ----------------------------------------------------------------- validate

def validate_pilot(sysml_path, jupyter):
    text = open(sysml_path, encoding="utf-8").read()
    nb = {"cells": [{"cell_type": "code", "metadata": {}, "outputs": [],
                     "execution_count": None, "source": text}],
          "metadata": {"kernelspec": {"name": "sysml", "language": "sysml",
                                      "display_name": "SysML"}},
          "nbformat": 4, "nbformat_minor": 5}
    with tempfile.TemporaryDirectory() as td:
        nbin = os.path.join(td, "check.ipynb")
        nbout = os.path.join(td, "check_out.ipynb")
        json.dump(nb, open(nbin, "w", encoding="utf-8"))
        cmd = [jupyter, "nbconvert", "--to", "notebook", "--execute",
               "--allow-errors", "--ExecutePreprocessor.kernel_name=sysml",
               "--ExecutePreprocessor.timeout=600",
               "--ExecutePreprocessor.startup_timeout=180",
               "--output", os.path.basename(nbout), "--output-dir", td, nbin]
        print("validating against SysML v2 Pilot kernel (first run ~1 min)...")
        r = subprocess.run(cmd, capture_output=True, text=True)
        if not os.path.exists(nbout):
            print("VALIDATION FAILED TO RUN:\n" + (r.stderr or r.stdout)[-2000:])
            return 2
        out = json.load(open(nbout, encoding="utf-8"))
        errors, warnings = [], []
        for cell in out.get("cells", []):
            for o in cell.get("outputs", []):
                stream = "".join(o.get("text", [])) if isinstance(
                    o.get("text"), list) else str(o.get("text", ""))
                for ln in stream.splitlines():
                    if ln.startswith("ERROR"):
                        errors.append(ln)
                    elif ln.startswith("WARNING"):
                        warnings.append(ln)
    print(f"validation: {len(errors)} error(s), {len(warnings)} warning(s)")
    for ln in errors[:40]:
        print("  " + ln)
    if len(errors) > 40:
        print(f"  ... and {len(errors) - 40} more")
    for ln in warnings[:15]:
        print("  " + ln)
    return 1 if errors else 0


# --------------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("tsv", help="SysML17_Migration_Export_w_Diagrams.vm output")
    ap.add_argument("-o", "--out", help="output .sysml path")
    ap.add_argument("--dialect", choices=["pilot", "msosa"], default="pilot")
    ap.add_argument("--report", help="output .xlsx migration report path")
    ap.add_argument("--no-report", action="store_true")
    ap.add_argument("--validate", action="store_true",
                    help="run the result through the Pilot kernel")
    ap.add_argument("--jupyter", default=JUPYTER_DEFAULT,
                    help="jupyter executable used by --validate "
                         "(default: first 'jupyter' on the PATH)")
    args = ap.parse_args(argv)

    nodes = load_tsv(args.tsv)
    mark_initial_transitions(nodes)
    roots = build_tree(nodes)
    if not roots:
        sys.exit("No root packages found in TSV")

    # QN -> element type map for reference-target checks; a Diagram can share
    # its owner's QN, so diagrams only fill gaps and never shadow model types
    et_by_qn = {}
    node_by_qn = {}
    for n in nodes:
        if n.et == "Diagram":
            et_by_qn.setdefault(n.qn, n.et)
            node_by_qn.setdefault(n.qn, n)
        else:
            et_by_qn[n.qn] = n.et
            node_by_qn[n.qn] = n
    em = Emitter(args.dialect, roots, et_by_qn=et_by_qn, node_by_qn=node_by_qn)
    em.w(0, f"// Generated by tsv_to_sysmlv2.py (dialect: {args.dialect}) "
            f"from {os.path.basename(args.tsv)}")
    em.w(0, "// TODO-MIGRATE comments mark constructs needing a human pass;")
    em.w(0, "// the companion .xlsx lists every non-migrated element/property/diagram.")
    for r in roots:
        em.visit(r, 0)

    # structural safety net: NOTHING may end the run unvisited. Any node
    # still 'pending' (e.g. owned by a construct whose handler does not
    # descend into it) becomes a report row instead of vanishing silently.
    stranded = [n for n in nodes if n.status == "pending"]
    for n in stranded:
        n.status = "notmig"
        n.extra["reason"] = ("not reached by any emitter (owned by a "
                             "construct that does not migrate children) - "
                             "migrate manually")
    if stranded:
        print(f"note: {len(stranded)} element(s) were unreachable by every "
              "handler - added to the report as not migrated")

    out = args.out or os.path.splitext(args.tsv)[0] + f"_{args.dialect}.sysml"
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(em.lines) + "\n")

    counts = {}
    for n in nodes:
        counts[n.status] = counts.get(n.status, 0) + 1
    print(f"wrote {out}")
    print(f"  elements: {counts.get('migrated', 0)} migrated, "
          f"{counts.get('notmig', 0)} TODO-stubbed (in report), "
          f"{counts.get('diagram', 0)} diagrams (in report), "
          f"{counts.get('consumed', 0)} consumed by parent constructs")

    if not args.no_report:
        project = last_seg(roots[0].qn) if roots else "model"
        rep = args.report or os.path.splitext(out)[0] + "_MigrationReport.xlsx"
        write_report(nodes, rep, project)
        print(f"wrote {rep}")

    if args.validate:
        sys.exit(validate_pilot(out, args.jupyter))


if __name__ == "__main__":
    main()
